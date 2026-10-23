import math
from io import BytesIO

from openpyxl.utils import get_column_letter
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.properties import PageSetupProperties

BULAN = ["JANUARI", "FEBRUARI", "MARET", "APRIL", "MEI", "JUNI", "JULI",
         "AGUSTUS", "SEPTEMBER", "OKTOBER", "NOVEMBER", "DESEMBER"]
SISI = Side(style="thin")
KOTAK = Border(left=SISI, right=SISI, top=SISI, bottom=SISI)
HEADER = ["NO. BERKAS", "NO. ITEM ARSIP", "KODE KLASIFIKASI", "URAIAN INFORMASI",
          "TANGGAL", "JUMLAH", "KETERANGAN", "KET. KLASIFIKASI KEAMANAN DAN AKSES ARSIP"]
LEBAR = [12, 13, 17, 70, 18, 11, 14, 26]
TENGAH = Alignment(horizontal="center", vertical="center", wrap_text=True)
ATAS = Alignment(horizontal="center", vertical="top", wrap_text=True)
KIRI = Alignment(horizontal="left", vertical="center", wrap_text=True)


def _isi_sheet(ws, item):
    """Isi satu tab. item: dict dengan kunci unit_label, bulan, tahun, grup, pengesah, gabungan."""
    gabungan = item["gabungan"]   
    header = HEADER + (["BAGIAN"] if gabungan else [])
    lebar = LEBAR + ([12] if gabungan else [])
    n = len(header)
    for i, w in enumerate(lebar, start=1):
        ws.column_dimensions[chr(64 + i)].width = w

    ws.merge_cells(f"A1:{chr(64 + n)}1")
    ws["A1"] = "DAFTAR ISI BERKAS"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A1"].alignment = Alignment(horizontal="center")

    info = [("UNIT KERJA", item["unit_label"]),
            ("TAHUN", item["tahun"]),
            ("BULAN", item["bulan"])]
    for baris, (label, isi) in enumerate(info, start=2):
        ws.cell(baris, 1, label)
        ws.cell(baris, 3, f": {isi}")

    for kol, teks in enumerate(header, start=1):
        c = ws.cell(5, kol, teks)
        c.font = Font(bold=True)
        c.alignment = TENGAH
        c.border = KOTAK
    ws.row_dimensions[5].height = 48

    gabung = []
    r = 6
    for g in item["grup"]:
        for k in range(1, n + 1):
            ws.cell(r, k).border = KOTAK
        ws.cell(r, 1, g["kategori"].nama or g["kategori"].kode)
        gabung.append((r, r, 1, n))
        r += 1
        for b in g["berkas"]:
            mulai = r
            for no_item, a in b["baris"]:
                for k in range(1, n + 1):
                    ws.cell(r, k).border = KOTAK
                ws.cell(r, 2, no_item).alignment = TENGAH
                ws.cell(r, 4, a.uraian).alignment = KIRI
                t = ws.cell(r, 5, a.tanggal)
                t.number_format = "d mmmm yyyy"
                t.alignment = TENGAH
                ws.cell(r, 6, a.jumlah_tampil).alignment = TENGAH
                ws.cell(r, 7, a.get_keterangan_display()).alignment = TENGAH
                ws.cell(r, 8, a.get_keamanan_display()).alignment = TENGAH
                if gabungan:
                    ws.cell(r, 9, a.unit_kerja.singkatan).alignment = TENGAH
                ws.row_dimensions[r].height = max(30, 15 * math.ceil(len(a.uraian) / 62))
                r += 1
            ws.cell(mulai, 1, b["no"]).alignment = ATAS
            ws.cell(mulai, 3, b["kode"]).alignment = ATAS
            if r - mulai > 1:
                gabung.append((mulai, r - 1, 1, 1))
                gabung.append((mulai, r - 1, 3, 3))

    for r1, r2, c1, c2 in gabung:
        ws.merge_cells(start_row=r1, end_row=r2, start_column=c1, end_column=c2)

    if not gabungan:
        pengesah = item["pengesah"]
        r += 1
        jabatan = f"{pengesah.jabatan}," if pengesah and pengesah.jabatan else "Kepala Bagian,"
        blok = [
            (r, "Mengetahui,", None),
            (r + 1, jabatan, None),
            (r + 5, pengesah.nama if pengesah and pengesah.nama else "Nama", "single"),
            (r + 6, f"NIP. {pengesah.nip}" if pengesah and pengesah.nip else "NIP.", None),
        ]
        for baris, teks, garis in blok:
            ws.merge_cells(start_row=baris, end_row=baris, start_column=7, end_column=8)
            c = ws.cell(baris, 7, teks)
            c.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
            if garis:
                c.font = Font(underline=garis)
        if len(jabatan) > 40:   
            ws.row_dimensions[r + 1].height = 30

    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)


def label_periode(awal, akhir):
    """Teks untuk isian BULAN dan TAHUN di bagian atas sheet."""
    if awal == akhir:
        bulan = BULAN[awal.month - 1]
    else:
        bulan = f"{BULAN[awal.month - 1]} - {BULAN[akhir.month - 1]}"
    tahun = str(awal.year) if awal.year == akhir.year else f"{awal.year} - {akhir.year}"
    return bulan, tahun

def buat_excel(lembar, rekap=None):
    """lembar: daftar dict (nama, unit_label, bulan, tahun, grup, pengesah, gabungan).
    rekap (opsional): dict untuk tab pertama berisi rekap total seluruh berkas."""
    wb = Workbook()
    wb.remove(wb.active)
    if rekap:
        _isi_rekap(wb.create_sheet(f"Rekap {rekap['tahun']}"), rekap)
    for item in lembar:
        ws = wb.create_sheet(item["nama"][:31])
        _isi_sheet(ws, item)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()

def buat_excel_rekap(judul, tahun, kolom, baris, jumlah, total_semua, grup=None, pakai_no=False):
    """Satu tab rekap. baris: list (nama bagian, [nilai per kolom]). grup: judul kelompok kolom (opsional)."""
    wb = Workbook()
    ws = wb.active
    ws.title = judul[:31]

    awal = ["NO", "BAGIAN"] if pakai_no else ["BAGIAN"]
    na = len(awal)
    n = na + len(kolom)
    huruf = get_column_letter

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n)
    ws["A1"] = judul.upper()
    ws["A1"].font = Font(bold=True, size=12)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws["A2"] = "TAHUN"
    ws["C2"] = f": {tahun}"

    r0 = 4
    if grup:
        for k, teks in enumerate(awal, start=1):
            ws.merge_cells(start_row=r0, start_column=k, end_row=r0 + 1, end_column=k)
            ws.cell(r0, k, teks)
        ws.merge_cells(start_row=r0, start_column=na + 1, end_row=r0, end_column=n)
        ws.cell(r0, na + 1, grup)
        for k, teks in enumerate(kolom, start=na + 1):
            ws.cell(r0 + 1, k, teks)
        r = r0 + 2
    else:
        for k, teks in enumerate(awal + list(kolom), start=1):
            ws.cell(r0, k, teks)
        r = r0 + 1
    for rr in range(r0, r):
        for k in range(1, n + 1):
            c = ws.cell(rr, k)
            c.font = Font(bold=True)
            c.alignment = TENGAH
            c.border = KOTAK
    ws.row_dimensions[r0].height = 32
    if grup:
        ws.row_dimensions[r0 + 1].height = 32

    for no, (nama, nilai) in enumerate(baris, start=1):
        if pakai_no:
            ws.cell(r, 1, no)
        ws.cell(r, na, nama)
        for k, v in enumerate(nilai, start=na + 1):
            ws.cell(r, k, v)
        for k in range(1, n + 1):
            ws.cell(r, k).border = KOTAK
            ws.cell(r, k).alignment = TENGAH
        r += 1

    if na > 1:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=na)
    ws.cell(r, 1, "Jumlah")
    for k, v in enumerate(jumlah, start=na + 1):
        ws.cell(r, k, v)
    r += 1
    if na > 1:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=na)
    ws.cell(r, 1, "Total")
    ws.merge_cells(start_row=r, start_column=na + 1, end_row=r, end_column=n)
    ws.cell(r, na + 1, total_semua)
    for rr in (r - 1, r):
        for k in range(1, n + 1):
            c = ws.cell(rr, k)
            c.font = Font(bold=True)
            c.alignment = TENGAH
            c.border = KOTAK

    ws.column_dimensions["A"].width = 8 if pakai_no else 14
    for k in range(2, n + 1):
        ws.column_dimensions[huruf(k)].width = 14 if (pakai_no and k == 2) else 16

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()