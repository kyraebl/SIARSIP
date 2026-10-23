import re
from collections import Counter
from datetime import date, datetime

from django.db import transaction
from openpyxl import load_workbook

from master.models import Kategori, KodeTujuan
from .models import Arsip

MAKS_BARIS = 5000

POLA_URAIAN = re.compile(
    r"^Surat\s+.+?\s+No\s*:\s*"
    r"(?P<kategori>[A-Za-z]{2})(?P<sub>\d*)\s*-\s*(?P<tujuan>[A-Za-z]+\d*)\s*/\s*(?P<nomor>\S+)"
    r"\s+Hal\s+(?P<perihal>.+)$",
    re.IGNORECASE | re.DOTALL,
)

POLA_BELUM_NOMOR = re.compile(
    r"^Surat\s+.+?\s+No\s*:\s*"
    r"(?P<kategori>[A-Za-z]{2})(?P<sub>\d*)\s*-\s*(?P<tujuan>[A-Za-z]+\d*)\s*/\s*(?:Hal\s+(?P<perihal>.+))?$",
    re.IGNORECASE | re.DOTALL,
)

POLA_KODE = re.compile(r"^(?P<kategori>[A-Za-z]{2})\s*(?P<sub>\d*)$")

POLA_NOMOR_KOSONG = re.compile(r"^[-._?xX]*$")

BULAN = {
    "januari": 1, "january": 1, "jan": 1,
    "februari": 2, "february": 2, "feb": 2,
    "maret": 3, "march": 3, "mar": 3,
    "april": 4, "apr": 4,
    "mei": 5, "may": 5,
    "juni": 6, "june": 6, "jun": 6,
    "juli": 7, "july": 7, "jul": 7,
    "agustus": 8, "august": 8, "agu": 8, "agt": 8, "aug": 8,
    "september": 9, "sep": 9, "sept": 9,
    "oktober": 10, "october": 10, "okt": 10, "oct": 10,
    "november": 11, "nov": 11,
    "desember": 12, "december": 12, "des": 12, "dec": 12,
}

class ImportGagal(Exception):
    """Kesalahan yang pesannya boleh ditampilkan langsung ke pengguna."""

def _teks(v):
    return " ".join(str(v).split()) if v is not None else ""

def _norm(v):
    return _teks(v).lower()

def _ambil(row, idx):
    if idx is None or idx >= len(row):
        return None
    return row[idx]

def _tanggal(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = _norm(v)
    if not s:
        return None
    try:
        m = re.match(r"^(\d{1,2})\s+([a-z]+)\.?\s+(\d{4})$", s)
        if m and m.group(2) in BULAN:
            return date(int(m.group(3)), BULAN[m.group(2)], int(m.group(1)))
        m = re.match(r"^(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{4})$", s)
        if m:
            return date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
        m = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})", s)
        if m:
            return date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    except ValueError:
        return None
    return None


def _tampil_tanggal(v):
    """Teks tanggal untuk tabel preview. Kalau tidak terbaca, tampilkan apa adanya."""
    d = _tanggal(v)
    return d.strftime("%d/%m/%Y") if d else _teks(v)


def _jumlah(v):
    m = re.search(r"\d+", _teks(v))
    return max(int(m.group()), 1) if m else 1


def _keterangan(v):
    s = _norm(v)
    if not s or s.startswith("asli"):
        return Arsip.Keterangan.ASLI
    if s.startswith(("copy", "copi", "salinan", "fotokopi")):
        return Arsip.Keterangan.COPY
    return None


def _keamanan(v):
    s = _norm(v)
    if not s:
        return Arsip.Keamanan.TERBUKA
    if "rahasia" in s:
        return Arsip.Keamanan.RAHASIA
    if "terbatas" in s:
        return Arsip.Keamanan.TERBATAS
    if any(k in s for k in ("biasa", "umum", "terbuka")):
        return Arsip.Keamanan.TERBUKA
    return None


def _cari_header(rows):
    """Cari baris judul kolom, lalu petakan nama kolom ke nomor kolomnya."""
    for i, row in enumerate(rows[:25]):
        teks = [_norm(c) for c in row]
        if "uraian informasi" in teks and any("tanggal" in t for t in teks):
            kolom = {}
            for j, t in enumerate(teks):
                if "kode klasifikasi" in t:
                    kolom.setdefault("kode", j)
                elif "uraian" in t:
                    kolom.setdefault("uraian", j)
                elif "tanggal" in t:
                    kolom.setdefault("tanggal", j)
                elif "jumlah" in t:
                    kolom.setdefault("jumlah", j)
                elif t == "keterangan":
                    kolom.setdefault("keterangan", j)
                elif "keamanan" in t:
                    kolom.setdefault("keamanan", j)
            return i, kolom
    return None, None


def proses_import(file, user, simpan=False, unit=None):
    """Baca file Excel Daftar Isi Berkas. simpan=False hanya mengecek tanpa menyimpan."""
    unit = unit or user.unit_kerja
    bebas = unit.uraian_bebas
    hasil = {"lembar": [], "dilewati": [], "total": 0, "siap": 0,
             "duplikat": 0, "galat": [], "belum_nomor": 0, "diimpor": 0,
             "baris": []}

    try:
        wb = load_workbook(file, data_only=True)
    except Exception:
        raise ImportGagal("File tidak bisa dibaca. Pastikan formatnya .xlsx, tidak rusak, dan tidak diberi password.")

    kategori = {k.kode.upper(): k for k in Kategori.objects.all()}
    tujuan = {t.kode.lower(): t for t in KodeTujuan.objects.all()}
    ada = set(Arsip.objects.filter(unit_kerja=unit).exclude(nomor_surat="").values_list(
        "kategori_id", "sub_kode", "kode_tujuan_id", "nomor_surat", "tanggal"))
    ada_kosong = {(k, sk, kt, t, _norm(pr)) for k, sk, kt, t, pr in Arsip.objects.filter(
        unit_kerja=unit, nomor_surat="").values_list(
        "kategori_id", "sub_kode", "kode_tujuan_id", "tanggal", "perihal")}
    ada_bebas = Counter((k, sk, t, _norm(pr)) for k, sk, t, pr in Arsip.objects.filter(
        unit_kerja=unit, kode_tujuan__isnull=True).values_list(
        "kategori_id", "sub_kode", "tanggal", "perihal"))
    lihat_bebas = Counter()
    baru = []

    def selesai(item, status, alasan="", kolom=()):
        """Simpan hasil cek satu baris untuk tabel preview."""
        item.update(status=status, alasan=alasan, kolom=list(kolom))
        hasil["baris"].append(item)
        if status == "galat":
            hasil["galat"].append({"lembar": item["lembar"], "baris": item["baris"],
                                   "uraian": item["uraian"][:90], "alasan": alasan})
        elif status == "belum":
            hasil["belum_nomor"] += 1
        elif status == "duplikat":
            hasil["duplikat"] += 1
        else:
            hasil["siap"] += 1

    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        idx, kol = _cari_header(rows)
        if idx is None:
            hasil["dilewati"].append(ws.title)
            continue

        jumlah_baris = 0
        kode_terakhir = ""
        for n, row in enumerate(rows[idx + 1:], start=idx + 2):
            uraian = _teks(_ambil(row, kol.get("uraian")))
            if not uraian:
                continue
            jumlah_baris += 1
            hasil["total"] += 1
            if hasil["total"] > MAKS_BARIS:
                raise ImportGagal(f"Terlalu banyak baris (maksimal {MAKS_BARIS} per file). Pecah file-nya dulu.")

            raw_tgl = _ambil(row, kol.get("tanggal"))
            item = {
                "lembar": ws.title, "baris": n, "uraian": uraian,
                "tanggal": _tampil_tanggal(raw_tgl),
                "jumlah": _teks(_ambil(row, kol.get("jumlah"))),
                "keterangan": _teks(_ambil(row, kol.get("keterangan"))),
                "keamanan": _teks(_ambil(row, kol.get("keamanan"))),
            }

            if bebas:
                item["kode"] = _teks(_ambil(row, kol.get("kode"))) or kode_terakhir
                kode_terakhir = item["kode"]
                mk = POLA_KODE.match(item["kode"])
                if not mk:
                    selesai(item, "galat", "Kode klasifikasi harus berbentuk seperti KU 03.", ["kode"])
                    continue
                kat = kategori.get(mk["kategori"].upper())
                if not kat:
                    selesai(item, "galat", f"Kategori {mk['kategori'].upper()} belum ada di data master.", ["kode"])
                    continue
                tgl = _tanggal(raw_tgl)
                if not tgl:
                    selesai(item, "galat", "Tanggal kosong atau formatnya tidak terbaca.", ["tanggal"])
                    continue
                ket = _keterangan(_ambil(row, kol.get("keterangan")))
                if ket is None:
                    selesai(item, "galat", "Kolom Keterangan harus Asli atau Copy.", ["keterangan"])
                    continue
                kea = _keamanan(_ambil(row, kol.get("keamanan")))
                if kea is None:
                    selesai(item, "galat", "Klasifikasi keamanan harus Biasa/Umum/Terbuka, Terbatas, atau Rahasia.", ["keamanan"])
                    continue
                kunci = (kat.id, mk["sub"], tgl, _norm(uraian))
                lihat_bebas[kunci] += 1
                if lihat_bebas[kunci] <= ada_bebas[kunci]:
                    selesai(item, "duplikat", "Sudah ada di sistem, tidak diimpor lagi.")
                    continue
                a = Arsip(
                    unit_kerja=unit, kode_tujuan=None, kategori=kat, sub_kode=mk["sub"],
                    nomor_surat="", perihal=uraian, tanggal=tgl,
                    jumlah=_jumlah(_ambil(row, kol.get("jumlah"))),
                    keterangan=ket, keamanan=kea, created_by=user,
                )
                a.uraian = a.rangkai_uraian()
                baru.append(a)
                selesai(item, "siap")
                continue

            belum = False
            m = POLA_URAIAN.match(uraian)
            if m and POLA_NOMOR_KOSONG.match(m["nomor"]):
                belum = True
            if not m:
                m = POLA_BELUM_NOMOR.match(uraian)
                belum = bool(m)
            if not m:
                selesai(item, "galat", "Format uraian tidak dikenali. Harus berbentuk: Surat ... No: KODE-Tujuan/Nomor Hal ...", ["uraian"])
                continue
            item["kode"] = f"{m['kategori'].upper()}{m['sub']}"
            kat = kategori.get(m["kategori"].upper())
            if not kat:
                selesai(item, "galat", f"Kategori {m['kategori'].upper()} belum ada di data master.", ["uraian"])
                continue
            tuj = tujuan.get(m["tujuan"].lower())
            if not tuj:
                selesai(item, "galat", f"Kode tujuan {m['tujuan']} belum ada di data master.", ["uraian"])
                continue
            tgl = _tanggal(raw_tgl)
            if not tgl:
                selesai(item, "galat", "Tanggal kosong atau formatnya tidak terbaca.", ["tanggal"])
                continue
            ket = _keterangan(_ambil(row, kol.get("keterangan")))
            if ket is None:
                selesai(item, "galat", "Kolom Keterangan harus Asli atau Copy.", ["keterangan"])
                continue
            kea = _keamanan(_ambil(row, kol.get("keamanan")))
            if kea is None:
                selesai(item, "galat", "Klasifikasi keamanan harus Biasa/Umum/Terbuka, Terbatas, atau Rahasia.", ["keamanan"])
                continue

            sub = m["sub"]
            nomor = "" if belum else m["nomor"].strip().upper()
            perihal = _teks(m["perihal"])
            if len(sub) > 6 or len(nomor) > 30 or not perihal:
                selesai(item, "galat", "Sub-kode, nomor surat, atau perihal tidak valid.", ["uraian"])
                continue

            if belum:
                kunci = (kat.id, sub, tuj.id, tgl, _norm(perihal))
                sudah, daftar = kunci in ada_kosong, ada_kosong
            else:
                kunci = (kat.id, sub, tuj.id, nomor, tgl)
                sudah, daftar = kunci in ada, ada
            if sudah:
                selesai(item, "duplikat", "Sudah ada di sistem, tidak diimpor lagi.")
                continue
            daftar.add(kunci)

            a = Arsip(
                unit_kerja=unit, kode_tujuan=tuj, kategori=kat, sub_kode=sub,
                nomor_surat=nomor, perihal=perihal, tanggal=tgl,
                jumlah=_jumlah(_ambil(row, kol.get("jumlah"))),
                keterangan=ket, keamanan=kea, created_by=user,
            )
            a.uraian = a.rangkai_uraian()
            baru.append(a)
            if belum:
                selesai(item, "belum", "Nomor surat belum keluar. Tetap diimpor, nomornya bisa dilengkapi nanti lewat tombol ubah.", ["uraian"])
            else:
                selesai(item, "siap")

        hasil["lembar"].append((ws.title, jumlah_baris))

    hasil["perlu_dicek"] = len(hasil["galat"]) + hasil["belum_nomor"]
    hasil["bisa_impor"] = hasil["siap"] + hasil["belum_nomor"]
    if simpan and baru:
        with transaction.atomic():
            Arsip.objects.bulk_create(baru, batch_size=500)
        hasil["diimpor"] = len(baru)
    return hasil